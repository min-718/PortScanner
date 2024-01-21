from fpdf import FPDF
from docx import Document

def download_results_to_pdf(results, file_name='port_details.pdf'):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", size=12)

    for result in results:
        pdf.cell(200, 10, txt=result, ln=True)

    pdf.output(file_name)

def download_results_to_txt(results, file_name='port_details.txt'):
    with open(file_name, 'w') as file:
        for result in results:
            file.write(result + '\n')

def download_results_to_word(results, file_name='port_details.docx'):
    doc = Document()

    for result in results:
        doc.add_paragraph(result)

    doc.save(file_name)